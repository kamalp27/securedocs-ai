"""DOCX Document Parser using python-docx."""

import re
from pathlib import Path
from typing import List, Optional

import docx

from app.rag.ingestion.base import (
    BaseDocumentParser,
    EmptyDocumentError,
    ExtractedSegment,
    MalformedDocumentError,
    ParsedDocument,
)


class DOCXParser(BaseDocumentParser):
    """Secure DOCX parser extracting paragraphs and table cells with section awareness."""

    def parse(self, file_path: Path, document_title: Optional[str] = None) -> ParsedDocument:
        title = document_title or file_path.stem
        segments: List[ExtractedSegment] = []
        current_offset = 0
        current_section = "General"

        try:
            doc = docx.Document(str(file_path))

            # 1. Parse paragraphs and headings
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                # Identify heading styles for section metadata
                style_name = p.style.name if p.style else ""
                if "heading" in style_name.lower() or "title" in style_name.lower():
                    current_section = text

                cleaned_text = re.sub(r"[ \t]+", " ", text)
                seg_start = current_offset
                seg_end = current_offset + len(cleaned_text)

                segments.append(
                    ExtractedSegment(
                        text=cleaned_text,
                        page_number=None,  # DOCX does not have native PDF page numbers
                        section=current_section,
                        start_char=seg_start,
                        end_char=seg_end,
                    )
                )
                current_offset = seg_end + 2

            # 2. Parse table contents
            for table_idx, table in enumerate(doc.tables):
                for row_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        row_text = " | ".join(row_cells)
                        seg_start = current_offset
                        seg_end = current_offset + len(row_text)
                        segments.append(
                            ExtractedSegment(
                                text=row_text,
                                page_number=None,
                                section=f"Table {table_idx + 1} Row {row_idx + 1}",
                                start_char=seg_start,
                                end_char=seg_end,
                            )
                        )
                        current_offset = seg_end + 2

        except EmptyDocumentError:
            raise
        except Exception as e:
            raise MalformedDocumentError(f"Failed to parse DOCX document: {e}") from e

        if not segments:
            raise EmptyDocumentError("DOCX document contains no readable text.")

        total_chars = sum(len(s.text) for s in segments)
        return ParsedDocument(title=title, segments=segments, total_chars=total_chars)
